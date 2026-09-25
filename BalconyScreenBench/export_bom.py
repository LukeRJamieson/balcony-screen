"""
export_bom.py
=============
Generates the Bill of Materials, cut list, cutting plan, fixings schedule and
drilling schedule from BalconyScreenBench.py, so the spreadsheet always matches
the Fusion model.

    python export_bom.py                 -> Balcony_Screen_BOM.xlsx (needs openpyxl)
    python export_bom.py --csv           -> CSV files instead (no dependencies)

Prices are Bunnings online prices checked 16 Sep 2026 - they are blue input
cells in the BOM sheet, so update them in the spreadsheet and totals follow.
"""

import csv
import os
import sys

HERE = os.path.dirname(os.path.realpath(__file__))
sys.path.insert(0, HERE)
import BalconyScreenBench as SL  # noqa: E402  (works without Fusion)

PRICE_NOTE = "Bunnings online price checked 16 Sep 2026"


def export_csv(model, out_dir):
    cl = SL.cut_list(model)
    with open(os.path.join(out_dir, "cut_list.csv"), "w", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(["Mark", "Description", "Material", "Section", "Length mm", "Qty", "Part IDs"])
        for r in cl:
            w.writerow([r["mark"], r["name"], SL.CATALOG[r["material"]]["desc"], r["section"],
                        r["length"], r["qty"], " ".join(r["ids"])])
    with open(os.path.join(out_dir, "cutting_plan.csv"), "w", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(["Material", "Stock #", "Stock mm", "Pieces (mark length)", "Offcut mm"])
        for b in SL.optimise_cuts(model):
            w.writerow([b["material"], b["n"], b["stock"],
                        " | ".join(f"{m} {l:g}" for m, l, _ in b["pieces"]), round(b["remaining"])])
    with open(os.path.join(out_dir, "fixings.csv"), "w", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(["FID", "Kind", "Note", "Head X", "Head Y", "Head Z", "Tip X", "Tip Y", "Tip Z"])
        for f in model.fasteners:
            w.writerow([f.fid, f.kind, f.note, *[round(v, 1) for v in f.head], *[round(v, 1) for v in f.tip]])
    with open(os.path.join(out_dir, "drilling.csv"), "w", newline="") as fh:
        rows = SL.drilling_schedule(model)
        w = csv.DictWriter(fh, fieldnames=list(rows[0].keys()))
        w.writeheader()
        w.writerows(rows)
    print("CSV files written to", out_dir)


def export_xlsx(model, path):
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    ARIAL = "Arial"
    f_norm = Font(name=ARIAL, size=10)
    f_bold = Font(name=ARIAL, size=10, bold=True)
    f_head = Font(name=ARIAL, size=10, bold=True, color="FFFFFF")
    f_title = Font(name=ARIAL, size=14, bold=True)
    f_input = Font(name=ARIAL, size=10, color="0000FF")
    f_link = Font(name=ARIAL, size=10, color="0563C1", underline="single")
    fill_head = PatternFill("solid", fgColor="0D5257")
    fill_input = PatternFill("solid", fgColor="FFFF00")
    fill_total = PatternFill("solid", fgColor="E7F0EF")
    thin = Side(style="thin", color="BFBFBF")
    border = Border(top=thin, bottom=thin, left=thin, right=thin)
    money = '$#,##0.00;($#,##0.00);-'

    wb = Workbook()

    def sheet(title, headers, widths):
        ws = wb.create_sheet(title)
        for i, (h, wdt) in enumerate(zip(headers, widths), start=1):
            c = ws.cell(row=1, column=i, value=h)
            c.font, c.fill, c.border = f_head, fill_head, border
            c.alignment = Alignment(wrap_text=True, vertical="center")
            ws.column_dimensions[get_column_letter(i)].width = wdt
        ws.freeze_panes = "A2"
        return ws

    cfg = model.cfg
    info = model.info

    # ---------------- Summary ----------------
    ws = wb.active
    ws.title = "Summary"
    ws.column_dimensions["A"].width = 44
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 70
    ws["A1"] = "Balcony lattice screen - 5 independent bench modules - BOM & cut list"
    ws["A1"].font = f_title
    ws["A2"] = "Generated from BalconyScreenBench.py - re-run export_bom.py after changing Config."
    ws["A2"].font = f_norm

    rows = [
        ("Key dimensions (mm)", None, None),
        ("Wall length (measured)", cfg.wall_length, "From Luke's measurement"),
        ("Wall height (approx.)", cfg.wall_height, "About 1 m - reference only, nothing is fixed to the wall"),
        
        ("Lattice panel layout", " / ".join(f"{w:g}" for w in cfg.panel_widths), "Symmetrical, joints land on posts"),
        ("Number of independent modules", info["n_modules"], "One lattice panel + its own bench each; not fixed together"),
        ("Gap between modules", cfg.module_gap, None),
        ("Overall length of the row", info["unit_length"], "Modules + gaps, centred on the wall"),
        ("Clearance each end", info["x_offset_each_end"], "Wall length minus unit length, halved"),
        ("Bench top height", cfg.bench_top_z, None),
        ("Lattice bottom above floor", cfg.lattice_z0, "Lifted clear of standing water, behind the bench"),
        ("Maximum height allowed", cfg.max_height, "Nothing projects above this"),
        ("Lower shelf height", cfg.shelf_top_z, None),
        ("Clear height under bench top", info["shelf_clear_height"], "Max pot height on lower shelf"),
        ("Bench depth from wall", info["bench_depth_from_wall"], None),
        ("Overall height (top of screen)", info["overall_height"], None),
        ("Screen height above wall top", info["height_above_wall"], None),
        ("Lattice thickness (assumed)", cfg.lattice_thickness, "NOT listed by Bunnings - measure a sheet"),
        (None, None, None),
        ("Inputs", None, None),
        ("Spare allowance - boxed screws", 0.10, "Extra screws for split/lost ones"),
        ("Spare allowance - loose items", 0.05, "Bolts, washers, packers"),
        (None, None, None),
        ("Estimated total (AUD, inc GST)", None, "Sum of BOM line totals"),
        (None, None, None),
        ("Stability - INDICATIVE ONLY (not an engineering check)", None, None),
        ("Total pot + soil load, all modules (kg)", cfg.design_pot_kg, "Edit: weigh or estimate your pots"),
        ("Module checked", "STAB_MOD", "Least stable module (modules are independent)"),
        ("Module share of pots (by width)", "STAB_SHARE", "Assumes pots spread evenly along the row"),
        ("Unit self weight (kg)", "STAB_SELF", "This module only, from timber volumes and densities"),
        ("Self-weight lever arm to front edge (m)", "STAB_SELF_LEVER", "Horizontal distance, centre of mass to front of legs"),
        ("Pot lever arm to front edge (m)", "STAB_POT_LEVER", "Pots assumed centred on bench top"),
        ("Drag coefficient on solid area", cfg.drag_coefficient, "Edit if you have better data"),
        ("Solid frontal area above wall top (m2)", "STAB_AREA", "20% of lattice + cover strips, above the wall (wall shelters the rest)"),
        ("Height of wind load centroid (m)", "STAB_Z", "Above deck"),
        ("Resisting moment (N.m)", "F_RES", "(self weight x lever + pot share x lever) x 9.81"),
        ("Estimated tipping gust (km/h)", "F_TIP", "Wind over the wall pushing the unit into the balcony"),
    ]
    st = min((SL.stability(model, cfg.design_pot_kg, mod) for mod in SL.module_names(model)),
             key=lambda d: d["tip_gust_kmh"])
    r0 = 4
    spare_screw_cell = spare_loose_cell = total_cell = None
    rowmap = {}
    for i, (a, _, _) in enumerate(rows):
        if a:
            rowmap[a] = r0 + i
    R = lambda label: f"B{rowmap[label]}"
    placeholders = {
        "STAB_MOD": f"{st['module']} ({st['width']:.0f} wide)",
        "STAB_SHARE": round(st["pot_share"], 4),
        "STAB_SELF": round(st["self_kg"], 1),
        "STAB_SELF_LEVER": round(st["self_lever_m"], 3),
        "STAB_POT_LEVER": round(st["pot_lever_m"], 3),
        "STAB_AREA": round(st["frontal_area_m2"], 3),
        "STAB_Z": round(st["wind_centroid_m"], 3),
        "F_RES": "=(" + R("Unit self weight (kg)") + "*" + R("Self-weight lever arm to front edge (m)") + "+"
                 + R("Total pot + soil load, all modules (kg)") + "*" + R("Module share of pots (by width)") + "*"
                 + R("Pot lever arm to front edge (m)") + ")*9.81",
        "F_TIP": "=SQRT(" + R("Resisting moment (N.m)") + "/(0.6*" + R("Drag coefficient on solid area") + "*"
                 + R("Solid frontal area above wall top (m2)") + "*" + R("Height of wind load centroid (m)") + "))*3.6",
    }
    rows = [(a, placeholders.get(b, b) if isinstance(b, str) else b, c) for a, b, c in rows]
    for i, (a, b, c) in enumerate(rows):
        r = r0 + i
        if a:
            ws.cell(row=r, column=1, value=a).font = f_bold if b is None and c is None or a == "Inputs" else f_norm
        if b is not None:
            cell = ws.cell(row=r, column=2, value=b)
            cell.font = f_norm
            cell.alignment = Alignment(horizontal="right")
        if c:
            ws.cell(row=r, column=3, value=c).font = f_norm
        if a == "Spare allowance - boxed screws":
            spare_screw_cell = f"Summary!$B${r}"
        if a == "Spare allowance - loose items":
            spare_loose_cell = f"Summary!$B${r}"
        if a and a.startswith("Spare"):
            ws.cell(row=r, column=2).font = f_input
            ws.cell(row=r, column=2).fill = fill_input
            ws.cell(row=r, column=2).number_format = "0%"
        if a and a.startswith("Estimated total"):
            total_cell = ws.cell(row=r, column=2)
            ws.cell(row=r, column=1).font = f_bold
        if a == "Module share of pots (by width)":
            ws.cell(row=r, column=2).number_format = "0.0%"
        if a in ("Total pot + soil load, all modules (kg)", "Drag coefficient on solid area"):
            ws.cell(row=r, column=2).font = f_input
            ws.cell(row=r, column=2).fill = fill_input
        if a == "Resisting moment (N.m)":
            ws.cell(row=r, column=2).number_format = "#,##0"
        if a == "Estimated tipping gust (km/h)":
            ws.cell(row=r, column=2).number_format = "0"
            ws.cell(row=r, column=1).font = f_bold
            ws.cell(row=r, column=2).font = f_bold
        if a and a.startswith("Stability"):
            ws.cell(row=r, column=1).font = f_bold
    notes_r = r0 + len(rows) + 1
    ws.cell(row=notes_r, column=1, value="Legend").font = f_bold
    ws.cell(row=notes_r + 1, column=1, value="Blue text on yellow = input you can edit").font = f_input
    ws.cell(row=notes_r + 1, column=1).fill = fill_input
    ws.cell(row=notes_r + 2, column=1,
            value="Prices: " + PRICE_NOTE + ". Packs (e.g. 25-pack bolts) are often cheaper per unit.").font = f_norm

    # ---------------- Parts ----------------
    wp = sheet("Parts", ["Part ID", "Module", "Mark", "Role", "Description", "Material key", "Length mm",
                         "X0", "X1", "Y0", "Y1", "Z0", "Z1"],
               [16, 8, 7, 9, 36, 12, 11, 9, 9, 9, 9, 9, 9])
    for i, p in enumerate(model.parts, start=2):
        b = p.box
        vals = [p.pid, p.module, p.mark, p.group, p.name, p.material, round(p.length, 1),
                round(b.x0, 1), round(b.x1, 1), round(b.y0, 1), round(b.y1, 1), round(b.z0, 1), round(b.z1, 1)]
        for j, v in enumerate(vals, start=1):
            c = wp.cell(row=i, column=j, value=v)
            c.font, c.border = f_norm, border
    parts_last = len(model.parts) + 1

    # ---------------- Cut List ----------------
    wc = sheet("Cut List", ["Mark", "Description", "Material", "Section", "Cut length (mm)", "Qty",
                            "Total length (m)", "Part IDs"],
               [7, 46, 44, 10, 14, 6, 14, 60])
    cl = SL.cut_list(model)
    for i, r in enumerate(cl, start=2):
        vals = [r["mark"], r["name"], SL.CATALOG[r["material"]]["desc"], r["section"], r["length"], r["qty"],
                f"=E{i}*F{i}/1000", " ".join(r["ids"])]
        for j, v in enumerate(vals, start=1):
            c = wc.cell(row=i, column=j, value=v)
            c.font, c.border = f_norm, border
        wc.cell(row=i, column=7).number_format = "0.00"
    last = len(cl) + 1
    wc.cell(row=last + 2, column=1,
            value="Lengths are finished sizes. Cut 90x45 and decking square; pre-drill all merbau (hardwood).").font = f_norm

    # ---------------- Cutting Plan ----------------
    wcp = sheet("Cutting Plan", ["Stock #", "Material key", "Material", "Stock length (mm)", "Pieces (mark: length)",
                                 "No. pieces", "Pieces total (mm)", "Kerf + trim (mm)", "Offcut (mm)",
                                 "Ask store to cut at (mm from one end)"],
                [8, 12, 44, 14, 60, 10, 14, 14, 12, 18])
    plans = SL.optimise_cuts(model)
    for i, b in enumerate(plans, start=2):
        total = sum(l for _, l, _ in b["pieces"])
        vals = [f"{b['material']}-{b['n']}", b["material"], SL.CATALOG[b["material"]]["desc"], b["stock"],
                " | ".join(f"{m}: {l:g}" for m, l, _ in b["pieces"]), len(b["pieces"]), round(total, 1),
                f"=F{i}*{cfg.kerf:g}+{cfg.end_trim:g}", f"=D{i}-G{i}-H{i}",
                b.get("store_cut_at") or ""]
        for j, v in enumerate(vals, start=1):
            c = wcp.cell(row=i, column=j, value=v)
            c.font, c.border = f_norm, border
    plan_last = len(plans) + 1
    wcp.cell(row=plan_last + 2, column=1,
             value=f"Assumes {cfg.kerf:g} mm saw kerf per cut and {cfg.end_trim:g} mm end trim per length. "
                   + (f"Every piece you carry is {cfg.max_carry_length / 1000:g} m or shorter; lengths sold longer "
                      f"get one in-store cut at the position shown."
                      if any(b.get("store_cut_at") for b in plans) else
                      "All lengths are bought whole and cut at home; no in-store cuts.")).font = f_norm

    # ---------------- Fixings ----------------
    wf = sheet("Fixings", ["FID", "Kind", "Joint", "Head X", "Head Y", "Head Z", "Tip X", "Tip Y", "Tip Z", "Length", "Module"],
               [34, 16, 44, 9, 9, 9, 9, 9, 9, 9, 8])
    for i, f in enumerate(model.fasteners, start=2):
        vals = [f.fid, f.kind, f.note, *[round(v, 1) for v in f.head], *[round(v, 1) for v in f.tip], round(f.length, 1), f.module]
        for j, v in enumerate(vals, start=1):
            c = wf.cell(row=i, column=j, value=v)
            c.font, c.border = f_norm, border
    fix_last = len(model.fasteners) + 1

    wk = sheet("Packers", ["Packer ID", "Where", "X0", "X1", "Y0", "Y1", "Z0", "Z1"], [26, 40, 9, 9, 9, 9, 9, 9])
    for i, k in enumerate(model.packers, start=2):
        b = k.box
        vals = [k.pid, k.note] + [round(v, 1) for v in (b.x0, b.x1, b.y0, b.y1, b.z0, b.z1)]
        for j, v in enumerate(vals, start=1):
            c = wk.cell(row=i, column=j, value=v)
            c.font, c.border = f_norm, border
    pk_last = len(model.packers) + 1

    # ---------------- Drilling ----------------
    wd = sheet("Drilling", ["Mark", "Part ID", "Fastener", "Fixing ID", "Along part (mm)", "Measured from",
                            "Across face (mm)", "Across measured from", "Drill direction", "Hole dia (mm)", "End grain"],
               [7, 11, 16, 30, 14, 12, 14, 18, 18, 11, 9])
    for i, r in enumerate(SL.drilling_schedule(model), start=2):
        vals = [r["mark"], r["part"], r["fastener"], r["fid"], r["along"], r["along_ref"], r["across"],
                r["across_ref"], r["drill_dir"], r["hole"] if r["hole"] else "", "yes" if r["end_grain"] else ""]
        for j, v in enumerate(vals, start=1):
            c = wd.cell(row=i, column=j, value=v)
            c.font, c.border = f_norm, border

    # ---------------- BOM ----------------
    wb_ = sheet("BOM", ["Item", "Category", "Bunnings product", "Buy unit", "Qty needed", "Pack size",
                        "Qty to buy", "Unit price (AUD)", "Line total (AUD)", "Link", "Notes",
                        "Alt supplier", "Alt buy unit", "Alt pack size", "Alt qty to buy",
                        "Alt unit price (AUD)", "Alt line total (AUD)", "Alt link", "Cheaper by (AUD)"],
                [5, 11, 50, 15, 11, 9, 10, 13, 14, 10, 50, 22, 16, 10, 11, 13, 14, 10, 12])
    # Alternative supply, checked 18 Sep 2026. Trade fastener suppliers sell the same items
    # by the box. Timber merchants near Mooroolbark don't publish prices, so those rows carry
    # a supplier and a note but no price - ring them with the quantities.
    ALT = {
        "BOLT_M10x120": dict(supplier="The Fastener Factory (Clayton / online)", unit="box of 50",
                             pack=50, price=38.09,
                             url="https://www.thefastenerfactory.com.au/m10-x-120mm-50pc-cup-head-bolt-nut-hot-dip-galvanised-as1390-class-4-6-uts",
                             note="Also Scrooz: trade box of 25 at $25.39 ($1.02 ea), free shipping over $99"),
        "WASHER_M10": dict(supplier="The Fastener Factory (Clayton / online)", unit="box of 500",
                           pack=500, price=22.87,
                           url="https://www.thefastenerfactory.com.au/m10-galvanised-flat-washers-10mm-x-22-5-x-2-0-500pc",
                           note="One box covers the job several times over; cheaper than 93 loose"),
        "T90x45": dict(supplier="Bowens Croydon / Mount Evelyn, Bayswater Mitre 10", unit="2.4 m length",
                       pack=1, price=None, url="https://bowens.com.au/stores/croydon",
                       note="Trade counter - price on application. Quote the whole timber list at once"),
        "D90x22": dict(supplier="Bowens Croydon / Mount Evelyn, Bayswater Mitre 10", unit="2.4 m length",
                       pack=1, price=None, url="https://bowens.com.au/stores/croydon",
                       note="Trade counter - price on application"),
        "M42x19": dict(supplier="Bowens Croydon / Mount Evelyn, Bayswater Mitre 10", unit="length",
                       pack=1, price=None, url="https://bowens.com.au/stores/croydon",
                       note="Merbau screening lengths vary by merchant - check stocked lengths"),
        "PACKER_10x72x100": dict(supplier="Trade counter (Bowens / Mitre 10)", unit="bag",
                                 pack=None, price=None, url="https://bowens.com.au/stores/croydon",
                                 note="Ask for plastic packers by the bag rather than singly"),
    }

    bom = []
    notes = {"T90x45": "Posts, bench frames, lattice rails (bottom / mid / top)",
             "D90x22": "Per module: 5 top + 4 shelf boards cut to module width",
             "M42x19": "Cover strips on the wall side (hardwood - pre-drill). Only sold in 2.7 m: "
                       "ask Bunnings to cut each length at the position on the Cutting Plan sheet"}
    used_stock = sorted({(b["material"], b["stock"]) for b in plans})
    for key in ("T90x45", "D90x22", "M42x19"):
        cat = SL.CATALOG[key]
        for mat, stock_len in used_stock:
            if mat != key:
                continue
            # one BOM row per stock length actually used in the cutting plan
            bom.append(dict(cat="Timber", key=key, desc=cat["desc"], unit=(f"{stock_len / 1000:g} m length" + (" (store cut)" if stock_len > cfg.max_carry_length else "")),
                            need=(f"=COUNTIFS('Cutting Plan'!$B$2:$B${plan_last},\"{key}\","
                                  f"'Cutting Plan'!$D$2:$D${plan_last},{stock_len})"),
                            pack=1, buy="=E{r}", price=cat["stock"][stock_len], url=cat["url"],
                            note=notes[key]))
    for key in ("LAT600", "LAT900"):
        cat = SL.CATALOG[key]
        bom.append(dict(cat="Lattice", key=key, desc=cat["desc"], unit="sheet",
                        need=f"=COUNTIF(Parts!$F$2:$F${parts_last},\"{key}\")", pack=1, buy="=E{r}",
                        price=cat["stock"][1], url=cat["url"], note="20% blockout, pine; no cutting needed"))
    for key, note in (("BOLT_M10x120", "Frame lap joints, 2 per joint. Far cheaper by the box - see alt supply"),
                      ("BOLT_M10x150", "The 900 module's middle frame only: four joints of 45 + 45 + 45"),
                      ("WASHER_M10", "One under each nut, both bolt lengths"),
                      ("SCREW_10Gx50", "Deck boards to bearers, 2 per board per bearer"),
                      ("SCREW_10Gx65", "Cover strips through lattice into posts and rails @ ~300 c/c"),
                      ("SCREW_10Gx100", "Lattice rails: down into bearers or through the end posts, 2 per rail end. Buy the 50-pack, not 8-packs"),
                      ("PACKER_10x72x100", "Feet under every leg (confirm sold each)")):
        fc = SL.FASTENER_CATALOG[key]
        if key == "WASHER_M10":
            # one washer per bolt, whichever length
            need = (f"=COUNTIF(Fixings!$B$2:$B${fix_last},\"BOLT_M10x120\")"
                    f"+COUNTIF(Fixings!$B$2:$B${fix_last},\"BOLT_M10x150\")")
        elif key == "PACKER_10x72x100":
            need = f"=COUNTA(Packers!$A$2:$A${pk_last})"
        else:
            need = f"=COUNTIF(Fixings!$B$2:$B${fix_last},\"{key}\")"
        spare = spare_screw_cell if fc["pack"] > 1 else spare_loose_cell
        bom.append(dict(cat="Fixings", key=key, desc=fc["desc"], unit="pack" if fc["pack"] > 1 else "each",
                        need=need, pack=fc["pack"], buy=f"=ROUNDUP(E{{r}}*(1+{spare})/F{{r}},0)",
                        price=fc["price"], url=fc["url"], note=note))

    for i, item in enumerate(bom, start=2):
        alt = ALT.get(item["key"])
        note = item["note"] + (". " + alt["note"] if alt and alt.get("note") else "")
        if alt and alt["price"] is not None:
            spare_ref = spare_screw_cell if alt["pack"] > 1 else spare_loose_cell
            alt_cells = [alt["supplier"], alt["unit"], alt["pack"],
                         f"=ROUNDUP(E{i}*(1+{spare_ref})/N{i},0)", alt["price"], f"=O{i}*P{i}",
                         "Supplier", f"=I{i}-Q{i}"]
        elif alt:
            alt_cells = [alt["supplier"], alt["unit"], "", "", "", "", "Supplier", ""]
        else:
            alt_cells = ["", "", "", "", "", "", "", ""]
        vals = [i - 1, item["cat"], item["desc"], item["unit"], item["need"], item["pack"],
                item["buy"].format(r=i), item["price"], f"=G{i}*H{i}", "Bunnings", note] + alt_cells
        for j, v in enumerate(vals, start=1):
            c = wb_.cell(row=i, column=j, value=v)
            c.font, c.border = f_norm, border
            c.alignment = Alignment(vertical="top", wrap_text=j in (3, 11))
        pc = wb_.cell(row=i, column=8)
        pc.font, pc.fill, pc.number_format = f_input, fill_input, money
        pc.comment = Comment(PRICE_NOTE, "BOM")
        wb_.cell(row=i, column=9).number_format = money
        lc = wb_.cell(row=i, column=10)
        lc.hyperlink, lc.font = item["url"], f_link
        for j in (12, 13, 14, 15, 16, 17, 18, 19):
            wb_.cell(row=i, column=j).alignment = Alignment(vertical="top", wrap_text=j in (12, 13))
        if alt:
            ap = wb_.cell(row=i, column=16)
            if alt["price"] is not None:
                ap.font, ap.fill, ap.number_format = f_input, fill_input, money
                ap.comment = Comment("Alternative supply checked 18 Sep 2026", "BOM")
            wb_.cell(row=i, column=17).number_format = money
            wb_.cell(row=i, column=19).number_format = money
            al = wb_.cell(row=i, column=18)
            al.hyperlink, al.font = alt["url"], f_link
    tr = len(bom) + 2
    wb_.cell(row=tr, column=8, value="TOTAL").font = f_bold
    tot = wb_.cell(row=tr, column=9, value=f"=SUM(I2:I{tr - 1})")
    tot.font, tot.number_format, tot.fill = f_bold, money, fill_total
    wb_.cell(row=tr, column=15, value="TOTAL, cheaper option per row").font = f_bold
    best = wb_.cell(row=tr, column=17, value=f"=SUM(I2:I{tr - 1})-SUM(S2:S{tr - 1})")
    best.font, best.number_format, best.fill = f_bold, money, fill_total
    wb_.cell(row=tr, column=18, value="Saving").font = f_bold
    sav = wb_.cell(row=tr, column=19, value=f"=SUM(S2:S{tr - 1})")
    sav.font, sav.number_format, sav.fill = f_bold, money, fill_total
    wb_.cell(row=tr + 1, column=12,
             value="Rows with an alt price are priced both ways; blank alt prices are trade counters that quote on request.").font = f_norm
    wb_.cell(row=tr + 2, column=3,
             value="Not included: exterior oil/stain, drill bits (11 mm auger, 3.5 mm pilot, countersink), "
                   "and any owners corporation fees.").font = f_norm

    total_cell.value = f"=BOM!I{tr}"
    total_cell.number_format = money
    total_cell.font = f_bold

    # Sheet order: Summary, BOM, Cut List, Cutting Plan, Drilling, Fixings, Packers, Parts
    order = ["Summary", "BOM", "Cut List", "Cutting Plan", "Drilling", "Fixings", "Packers", "Parts"]
    wb._sheets = [wb[n] for n in order]
    wb.save(path)
    print("Saved", path)


if __name__ == "__main__":
    model = SL.build_model()
    if model.warnings:
        print("WARNINGS:\n  " + "\n  ".join(model.warnings))
    if "--csv" in sys.argv:
        export_csv(model, HERE)
    else:
        out = sys.argv[1] if len(sys.argv) > 1 and not sys.argv[1].startswith("-") else os.path.join(HERE, "Balcony_Screen_BOM.xlsx")
        export_xlsx(model, out)
